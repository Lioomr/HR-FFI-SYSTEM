import json
import re
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import openpyxl
import pytest
from attendance.models import AttendanceRecord, BioTimeEmployeeMap
from audit.models import AuditLog
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.cache import cache
from django.test import RequestFactory
from django.urls import resolve
from django.utils import timezone
from employees.models import EmployeeProfile
from employees.services.archiving import retire_biotime_mapping_and_archive_profile
from leaves.models import LeaveRequest, LeaveType
from loans.models import LoanRequest
from organization.models import OrganizationNode, UserOrganizationAccess
from penalties.models import PenaltyCatalog, PenaltyDeduction, PenaltyRecord
from pypdf import PdfReader
from rents.models import Rent, RentType
from rest_framework.test import APIClient

from . import dto, engine
from .i18n import enum_directory
from .models import (
    AccommodationAssignment,
    PayrollAdjustment,
    PayrollCarryForward,
    PayrollRun,
    PayrollSettings,
    Payslip,
)
from .settings_serializers import PayrollSettingsSerializer

pytestmark = pytest.mark.django_db


@pytest.fixture
def ctx():
    # SQLite reuses PKs after rollback; keep throttle/cache state per scenario.
    cache.clear()
    company = OrganizationNode.objects.create(code="P2_A", name="Phase 2", node_type="company")
    other = OrganizationNode.objects.create(code="P2_B", name="Other company", node_type="company")
    user = get_user_model().objects.create_user(email="p2hr@test.com", password="password")
    user.groups.add(Group.objects.get_or_create(name="HRManager")[0])
    UserOrganizationAccess.objects.create(user=user, organization=company)
    employee = get_user_model().objects.create_user(email="p2employee@test.com", password="password")
    employee.groups.add(Group.objects.get_or_create(name="Employee")[0])
    profile = EmployeeProfile.objects.create(
        company=company,
        user=employee,
        employee_id="P2-001",
        hire_date=date(2024, 1, 1),
        basic_salary="6000.00",
        accommodation_allowance="3000.00",
        transportation_allowance="1500.00",
        telephone_allowance="500.00",
        total_salary="11000.00",
        allowed_overtime=20,
    )
    BioTimeEmployeeMap.objects.create(
        employee_profile=profile, biotime_emp_code="P2-001", attendance_start_date=date(2024, 1, 1)
    )
    settings = PayrollSettings.objects.create(company=company, no_medical_insurance=True)
    AccommodationAssignment.objects.create(
        company=company,
        employee_profile=profile,
        status="none",
        created_by=user,
        updated_by=user,
        start_date=date(2020, 1, 1),
        end_date=date(2030, 12, 31),
    )
    client = APIClient()
    client.force_authenticate(user)
    request = RequestFactory().post("/payroll-runs/", HTTP_X_ACTIVE_COMPANY_ID=str(company.pk))
    request.user = user
    return SimpleNamespace(
        company=company,
        other=other,
        user=user,
        employee=employee,
        profile=profile,
        settings=settings,
        client=client,
        headers={"HTTP_X_ACTIVE_COMPANY_ID": str(company.pk)},
        request=request,
    )


def create(ctx, month=9):
    response = ctx.client.post("/payroll-runs/", {"year": 2026, "month": month}, format="json", **ctx.headers)
    assert response.status_code == 201, response.data
    return PayrollRun.objects.get(pk=response.data["data"]["id"])


def edit(ctx, run, changes):
    run.refresh_from_db()
    item = run.items.filter(is_active=True, employee_profile=ctx.profile).get()
    response = ctx.client.patch(
        f"/payroll-runs/{run.pk}/items/{item.pk}/",
        {"expected_revision": run.revision, "changes": changes, "reason": "Approved correction"},
        format="json",
        **ctx.headers,
    )
    assert response.status_code == 200, response.data
    return response.data["data"]


def finalize(ctx, run, status=200):
    response = ctx.client.post(f"/payroll-runs/{run.pk}/finalize/", {"confirm": True}, format="json", **ctx.headers)
    assert response.status_code == status, response.data
    run.refresh_from_db()
    return response


def test_create_accountless_totals_draft_visibility_and_noop_sync(ctx):
    accountless = EmployeeProfile.objects.create(
        company=ctx.company,
        employee_id="P2-002",
        hire_date=date(2024, 1, 1),
        basic_salary="1000.00",
        total_salary="1000.00",
    )
    run = create(ctx)
    assert run.total_employees == 2 and run.total_gross == Decimal("12000")
    assert Payslip.objects.get(payroll_run=run, employee_profile=accountless).employee_id is None
    assert set(run.payslips.values_list("status", flat=True)) == {"DRAFT"}
    before = AuditLog.objects.count()
    result = engine.recalculate(run)
    assert result["changed_items"] == 0 and AuditLog.objects.count() == before
    assert dto.reconciliation(run)["matches"]
    ctx.client.force_authenticate(ctx.employee)
    assert ctx.client.get("/employee/payslips/", **ctx.headers).data["data"]["items"] == []


@pytest.mark.parametrize(
    "limit,entered,paid,code",
    [
        (20, "24.00", "20.00", "overtime_over_allowed"),
        (0, "4.00", "0.00", "overtime_over_allowed"),
        (None, "4.00", "4.00", "overtime_limit_not_set"),
    ],
)
def test_overtime_month_cap_and_override(ctx, limit, entered, paid, code):
    ctx.profile.allowed_overtime = limit
    ctx.profile.save(update_fields=["allowed_overtime"])
    run = create(ctx)
    item = edit(ctx, run, {"overtime_normal_hours": entered, "weekly_rest_work_days": "1.00"})
    assert item["overtime"]["payable_normal_hours"] == paid
    assert item["weekly_rest_work"] == "300.00"
    assert code in {w["code"] for w in item["warnings"]}
    item = edit(ctx, run, {"override_overtime_eligibility": True})
    assert item["overtime"]["payable_normal_hours"] == entered
    assert AuditLog.objects.filter(action="payroll_overtime_override", actor=ctx.user).count() == 1


def test_mapped_absence_unpaid_leave_paid_leave_refresh(ctx):
    AttendanceRecord.objects.create(employee_profile=ctx.profile, date=date(2026, 9, 1), status="ABSENT")
    unpaid = LeaveType.objects.create(company=ctx.company, name="Unpaid", code="UP", is_paid=False)
    paid = LeaveType.objects.create(company=ctx.company, name="Paid", code="PD", is_paid=True)
    LeaveRequest.objects.create(
        company=ctx.company,
        employee_profile=ctx.profile,
        leave_type=unpaid,
        start_date=date(2026, 9, 2),
        end_date=date(2026, 9, 2),
        status="approved",
    )
    run = create(ctx)
    assert dto.item_dto(run.items.get())["absence_deduction"] == "733.00"
    LeaveRequest.objects.create(
        company=ctx.company,
        employee_profile=ctx.profile,
        leave_type=paid,
        start_date=date(2026, 9, 1),
        end_date=date(2026, 9, 1),
        status="approved",
    )
    engine.recalculate(run, request=ctx.request)
    assert dto.item_dto(run.items.get())["absence_deduction"] == "367.00"
    assert PayrollAdjustment.objects.filter(run=run, type="LEAVE_INPUT_CHANGE").exists()


def test_unmapped_zero_rows_warning_override_and_strict_finalize(ctx):
    ctx.profile.biotime_mapping.all().delete()
    ctx.settings.require_manual_absence_confirmation = True
    ctx.settings.save()
    run = create(ctx)
    item = dto.item_dto(run.items.get())
    assert item["absence_deduction"] == "0.00" and not item["absence"]["source_available"]
    assert "attendance_source_unavailable" in {w["code"] for w in item["warnings"]}
    finalize(ctx, run, 422)
    assert run.status == "DRAFT"
    item = edit(ctx, run, {"absence_days_override": "3.00"})
    assert item["absence_deduction"] == "1100.00"
    assert item["absence"]["manual_confirmed"]
    finalize(ctx, run)


def test_missing_hire_blocks_finalize_until_profile_corrected(ctx):
    ctx.profile.hire_date = None
    ctx.profile.save(update_fields=["hire_date"])
    run = create(ctx)
    response = finalize(ctx, run, 422)
    assert "missing_hire_date" in {e["code"] for e in response.data["errors"]}
    assert run.payslips.get().status == "DRAFT"
    ctx.profile.hire_date = date(2024, 1, 1)
    ctx.profile.save(update_fields=["hire_date"])
    finalize(ctx, run)


def test_settings_snapshot_refresh_and_history_shapes(ctx):
    run = create(ctx)
    response = ctx.client.patch(
        "/payroll-settings/",
        {
            "expected_revision": 1,
            "provision_proration": "PRORATED_BY_DAYS",
            "annual_air_ticket": "2400.00",
            "reason": "Annual policy",
        },
        format="json",
        **ctx.headers,
    )
    assert response.status_code == 200, response.data
    engine.recalculate(run)
    run.refresh_from_db()
    assert run.settings_snapshot["provision_proration"] == "FULL_MONTH"
    response = ctx.client.post(
        f"/payroll-runs/{run.pk}/recalculate/",
        {"expected_revision": run.revision, "refresh_settings": True, "reason": "Adopt latest policy"},
        format="json",
        **ctx.headers,
    )
    assert response.status_code == 200, response.data
    run.refresh_from_db()
    assert run.settings_snapshot["provision_proration"] == "PRORATED_BY_DAYS"
    assert PayrollAdjustment.objects.filter(run=run, type="SETTINGS_REFRESH").count() == 1
    value = ctx.client.get("/payroll-settings/history/", **ctx.headers).data["data"]["items"][0]
    assert value["settings_revision"] == 2 and "run_revision" not in value


def loan(ctx, amount):
    return LoanRequest.objects.create(
        company=ctx.company,
        employee=ctx.employee,
        employee_profile=ctx.profile,
        requested_amount=amount,
        approved_amount=amount,
        status="approved",
        target_deduction_year=2026,
        target_deduction_month=9,
    )


def test_loan_pending_priority_cancel_restores_and_replacement(ctx):
    row = loan(ctx, "11000.00")
    AttendanceRecord.objects.create(employee_profile=ctx.profile, date=date(2026, 9, 1), status="ABSENT")
    run = create(ctx)
    row.refresh_from_db()
    assert row.status == "approved" and row.deduction_payroll_run_id is None
    assert "claim_unaffordable" in {w["code"] for w in dto.validation(run)["warnings"]}
    edit(ctx, run, {"absence_days_override": "0.00"})
    row.refresh_from_db()
    assert row.status == "deducted" and row.deduction_payroll_run_id == run.pk
    run.refresh_from_db()
    response = ctx.client.post(
        f"/payroll-runs/{run.pk}/cancel/",
        {"expected_revision": run.revision, "reason": "Replace incorrect inputs"},
        format="json",
        **ctx.headers,
    )
    assert response.status_code == 200, response.data
    row.refresh_from_db()
    assert (
        row.status == "approved"
        and row.deduction_payroll_run_id is None
        and row.deducted_at is None
        and row.deducted_amount is None
    )
    assert not run.payslips.filter(is_active=True).exists()
    run.refresh_from_db()
    assert dto.reconciliation(run)["matches"]
    assert AuditLog.objects.filter(action="payroll_loan_restored", entity_id=str(row.pk)).count() == 1
    assert (
        ctx.client.post(
            f"/payroll-runs/{run.pk}/cancel/",
            {"reason": "Retry", "expected_revision": run.revision},
            format="json",
            **ctx.headers,
        ).status_code
        == 200
    )
    replacement = create(ctx)
    assert replacement.pk != run.pk and PayrollRun.objects.filter(company=ctx.company, month=9).count() == 2
    assert (
        ctx.client.post("/payroll-runs/", {"year": 2026, "month": 9}, format="json", **ctx.headers).status_code == 409
    )


def test_carry_forward_finalization_reservation_release_and_collect_once(ctx):
    ctx.profile.deduct_gosi_from_employee = True
    ctx.profile.save(update_fields=["deduct_gosi_from_employee"])
    run = create(ctx)
    item = edit(ctx, run, {"absence_days_override": "30.00", "zero_salary_reason": "full_absence"})
    assert item["net_salary"] == "0.00" and item["gosi_company_provision"] == "0.00"
    finalize(ctx, run)
    balance = PayrollCarryForward.objects.get(origin_run=run)
    assert balance.type == "GOSI_EMPLOYEE" and balance.remaining_amount == Decimal("180")
    next_run = create(ctx, 10)
    balance.refresh_from_db()
    assert balance.remaining_amount == Decimal("180")
    assert dto.item_dto(next_run.items.get())["gosi_employee_deduction"] == "360.00"
    finalize(ctx, next_run)
    balance.refresh_from_db()
    assert balance.remaining_amount == 0 and balance.status == "SETTLED"
    finalize(ctx, next_run)
    assert balance.remaining_amount == 0


def test_finalization_reconciliation_repeated_publish_and_employee_whitelist(ctx):
    run = create(ctx)
    edit(ctx, run, {"absence_days_override": "2.00", "bonus": "50.00", "other_deduction": "20.00"})
    finalize(ctx, run)
    assert run.status == "COMPLETED" and run.payslips.get().status == "COMPLETED"
    finalize(ctx, run)
    assert AuditLog.objects.filter(action="payroll_run_finalized", entity_id=str(run.pk)).count() == 1
    with patch("in_app_notifications.dispatcher.dispatch_notification_channels") as dispatch:
        for _ in range(2):
            response = ctx.client.post(f"/payroll-runs/{run.pk}/generate-payslips/", {}, format="json", **ctx.headers)
            assert response.status_code == 200, response.data
    run.refresh_from_db()
    assert run.status == "PAID" and run.payslips.get().status == "PAID"
    assert AuditLog.objects.filter(action="payslips_generated", entity_id=str(run.pk)).count() == 1
    assert dispatch.call_count == 0  # Test transaction has not committed.
    ctx.client.force_authenticate(ctx.employee)
    slip = run.payslips.get()
    value = ctx.client.get(f"/employee/payslips/{slip.pk}/", **ctx.headers).data["data"]
    assert {d["type"] for d in value["deductions"]} == {"ABSENCE", "OTHER_DEDUCTION"}
    assert (
        not {
            "provisions",
            "settings_snapshot",
            "total_company_cost",
            "cost_incomplete",
            "monthly_salary_components",
            "salary_adjusted",
            "manual_inputs",
        }
        & value.keys()
    )
    assert not any({"id", "source_id", "source_kind", "requested_amount"} & d.keys() for d in value["deductions"])
    pdf = ctx.client.get(f"/employee/payslips/{slip.pk}/download/", **ctx.headers)
    assert pdf.status_code == 200
    text = " ".join(page.extract_text() for page in PdfReader(BytesIO(pdf.content)).pages)
    assert "absence" in text.lower() and "other deduction" in text.lower() and "733.00" in text
    assert "Company cost" not in text and "Medical" not in text and "Provision" not in text


def test_scope_permissions_stale_error_native_ints_and_no_recalc_mutation(ctx):
    run = create(ctx)
    original = dto.run_dto(run)
    response = ctx.client.post(
        f"/payroll-runs/{run.pk}/recalculate/",
        {"expected_revision": run.revision + 1, "reason": "Correction"},
        format="json",
        **ctx.headers,
    )
    assert response.status_code == 409
    value = response.data["errors"][0]
    assert value["code"] == "stale_revision" and isinstance(value["params"]["current_revision"], int)
    run.refresh_from_db()
    assert dto.run_dto(run) == original
    foreign = PayrollRun.objects.create(company=ctx.other, year=2026, month=9)
    for suffix in ("", "items/", "summary/", "validation-report/", "history/", "accommodation-report/", "export/"):
        assert ctx.client.get(f"/payroll-runs/{foreign.pk}/{suffix}", **ctx.headers).status_code == 404
    ctx.client.force_authenticate(ctx.employee)
    for path in ("/payroll-runs/", "/payroll-settings/", "/payroll-accommodation-assignments/"):
        assert ctx.client.get(path, **ctx.headers).status_code == 403


@pytest.mark.parametrize(
    "changes",
    [
        {"bonus": 50},
        {"bonus": "1e2"},
        {"overtime_normal_hours": "1e2"},
        {"overtime_normal_hours": 2},
        {"net_salary": "100.00"},
    ],
)
def test_strict_write_validation_rolls_back(ctx, changes):
    run = create(ctx)
    item = run.items.get()
    before = AuditLog.objects.count()
    response = ctx.client.patch(
        f"/payroll-runs/{run.pk}/items/{item.pk}/",
        {"expected_revision": run.revision, "changes": changes, "reason": "Correction"},
        format="json",
        **ctx.headers,
    )
    assert response.status_code == 422, response.data
    run.refresh_from_db()
    assert run.revision == 1 and AuditLog.objects.count() == before


def test_audit_failure_rolls_back_creation_and_loan_deduction(ctx):
    row = loan(ctx, "500.00")
    with patch("payroll.engine.audit", side_effect=RuntimeError("audit unavailable")):
        response = ctx.client.post("/payroll-runs/", {"year": 2026, "month": 9}, format="json", **ctx.headers)
    assert response.status_code == 500
    assert not PayrollRun.objects.filter(company=ctx.company).exists()
    row.refresh_from_db()
    assert row.status == "approved" and row.deduction_payroll_run_id is None


def test_rent_non_payroll_resident_rounding_report_frozen_and_exports(ctx):
    ctx.profile.payroll_accommodation_assignments.all().delete()
    type_ = RentType.objects.create(company=ctx.company, code="HOUSING", name_en="Housing")
    rent = Rent.objects.create(
        company=ctx.company,
        rent_type=type_,
        property_name_en="Shared home",
        annual_rent_value="12000.00",
        recurrence="MONTHLY",
        amount="1000.00",
        start_date=date(2026, 1, 1),
        due_day=1,
    )
    others = [
        EmployeeProfile.objects.create(
            company=ctx.company,
            employee_id=f"P2-R{i}",
            hire_date=date(2024, 1, 1),
            basic_salary="1000.00",
            total_salary="1000.00",
            employment_status="PREHIRE" if i == 2 else "ACTIVE",
        )
        for i in (1, 2)
    ]
    for profile in [ctx.profile, *others]:
        AccommodationAssignment.objects.create(
            company=ctx.company,
            employee_profile=profile,
            status="assigned",
            rent=rent,
            start_date=date(2026, 9, 1),
            end_date=date(2026, 9, 30),
        )
    run = create(ctx)
    allocation = run.rent_allocations.get()
    assert allocation.rounding_difference["numerator"] == "1.00"

    def legacy_numerators(value):
        if isinstance(value, dict):
            return {
                key: item.removesuffix(".00") if key == "numerator" else legacy_numerators(item)
                for key, item in value.items()
            }
        if isinstance(value, list):
            return [legacy_numerators(item) for item in value]
        return value

    stored_fields = (
        "monthly_rent",
        "resident_allocations",
        "slot_inputs",
        "vacant_unallocated_raw",
        "not_in_payroll_raw",
        "rounding_difference",
    )
    # Old snapshots retain integer-string numerators; normalize only their
    # outward representation and continue reconciling them exactly.
    for key in stored_fields:
        setattr(allocation, key, legacy_numerators(getattr(allocation, key)))
    allocation.save(update_fields=stored_fields)
    before = {key: getattr(allocation, key) for key in stored_fields}

    def assert_exact_share_format(value):
        if isinstance(value, dict):
            if {"numerator", "denominator", "decimal"} <= value.keys():
                assert isinstance(value["numerator"], str) and re.fullmatch(r"-?\d+\.\d{2}", value["numerator"])
            for item in value.values():
                assert_exact_share_format(item)
        elif isinstance(value, list):
            for item in value:
                assert_exact_share_format(item)

    response = ctx.client.get(f"/payroll-runs/{run.pk}/accommodation-report/", **ctx.headers)
    assert response.status_code == 200, response.data
    value = response.data["data"]["items"][0]
    assert_exact_share_format(value)
    assert [r["rounded_share"] for r in value["residents"]] == ["333.00"] * 3
    assert value["rounding_difference"]["decimal"] == "1" and value["in_payroll_rounded"] == "666.00"
    assert value["rounding_difference"]["numerator"] == "1.00"
    assert value["raw_reconciliation"]["matches"]
    response = ctx.client.get(f"/payroll-runs/{run.pk}/summary/", **ctx.headers)
    assert response.status_code == 200, response.data
    assert_exact_share_format(response.data)
    assert response.data["data"]["rent_rounding_difference"]["numerator"] == "1.00"
    rent.property_name_en = "Changed live property"
    rent.save(update_fields=["property_name_en"])
    assert dto.rent_dto(run.rent_allocations.get())["rent_name"] == "Shared home"
    response = ctx.client.get(f"/payroll-runs/{run.pk}/export/", {"file_format": "xlsx"}, **ctx.headers)
    assert response.status_code == 200
    wb = openpyxl.load_workbook(BytesIO(response.content))
    assert wb.sheetnames == ["Payroll", "Lines", "Provisions", "Validation", "Adjustments", "Accommodation"]
    assert wb["Payroll"].cell(2, 5).value == "6000.00"
    assert wb["Accommodation"].cell(2, 9).value == 3
    for row in wb["Accommodation"].iter_rows(min_row=2, values_only=True):
        assert isinstance(row[7], str) and re.fullmatch(r"-?\d+\.\d{2}", row[7])
    assert response["Cache-Control"] == "private, no-store"
    response = ctx.client.get(f"/payroll-runs/{run.pk}/export/", {"file_format": "pdf"}, **ctx.headers)
    assert response.status_code == 200
    text = " ".join(page.extract_text() for page in PdfReader(BytesIO(response.content)).pages)
    assert "rounding_difference: 1.00/1" in text
    allocation.refresh_from_db()
    assert {key: getattr(allocation, key) for key in stored_fields} == before


def test_approved_contract_dto_fields_and_nested_shapes(ctx):
    examples = json.loads(Path(__file__).with_name("contract_examples.json").read_text(encoding="utf-8"))
    assert sum(len(values) for values in enum_directory().values()) == 124
    assert len(enum_directory()["validation_codes"]) == 44
    run = create(ctx)
    edit(ctx, run, {"absence_days_override": "2.00", "bonus": "50.00"})
    run.refresh_from_db()

    def shape(actual, expected, path=""):
        if isinstance(expected, dict):
            assert set(actual) == set(expected), (path, set(actual) - set(expected), set(expected) - set(actual))
            for key in expected:
                shape(actual[key], expected[key], path + "." + key)
        elif isinstance(expected, list) and actual and expected:
            for element in actual:
                shape(element, expected[0], path + "[]")

    for name, value in {
        "enums": enum_directory(),
        "settings": PayrollSettingsSerializer(ctx.settings).data,
        "run": dto.run_dto(run),
        "item": dto.item_dto(run.items.get()),
        "validation": dto.validation(run),
        "assignment": dto.assignment_dto(ctx.profile.payroll_accommodation_assignments.get()),
        "employee": dto.employee_dto(run.payslips.get()),
        "summary": dto.summary(run),
    }.items():
        shape(value, examples[name], name)

    # Earlier frozen snapshots may retain the retired assumption. Hide only
    # that response entry without rewriting historical evidence or money.
    item = run.items.get()
    original_assumptions = dto.item_dto(item)["assumptions"]
    removed = "payroll.assumption.provision_proration_pending"
    item.input_snapshot["result"]["assumptions"].append(removed)
    item.save(update_fields=["input_snapshot"])
    assert dto.item_dto(item)["assumptions"] == original_assumptions
    item.refresh_from_db()
    assert removed in item.input_snapshot["result"]["assumptions"]


def test_gosi_config_and_noop_history(ctx):
    path = f"/payroll-settings/employees/{ctx.profile.pk}/"
    response = ctx.client.patch(
        path,
        {"expected_revision": 0, "deduct_gosi_from_employee": True, "reason": "Arrangement"},
        format="json",
        **ctx.headers,
    )
    assert response.status_code == 200, response.data
    assert response.data["data"]["settings_revision"] == 1
    count = AuditLog.objects.count()
    assert (
        ctx.client.patch(
            path,
            {"expected_revision": 1, "deduct_gosi_from_employee": True, "reason": "No change"},
            format="json",
            **ctx.headers,
        ).status_code
        == 200
    )
    assert AuditLog.objects.count() == count
    history = ctx.client.get(path + "history/", **ctx.headers).data["data"]["items"][0]
    assert history["type"] == "EMPLOYEE_GOSI_CHANGE" and history["settings_revision"] == 1
    run = create(ctx)
    item = dto.item_dto(run.items.get())
    assert item["gosi_company_provision"] == "0.00" and item["gosi_employee_deduction"] == "180.00"


def test_excluded_after_lifecycle_correction_retains_item_history(ctx):
    run = create(ctx)
    edit(ctx, run, {"bonus": "50.00"})
    item_id = run.items.get().pk
    ctx.profile.hire_date = date(2026, 10, 1)
    ctx.profile.save(update_fields=["hire_date"])
    engine.recalculate(run, request=ctx.request, reason="Corrected joining date", advance_revision=True)
    run.refresh_from_db()
    assert run.total_employees == 0 and run.total_net == 0
    assert not run.items.get(pk=item_id).is_active and not run.payslips.filter(is_active=True).exists()
    assert PayrollAdjustment.objects.filter(item_id=item_id).exists()
    assert dto.reconciliation(run)["matches"]


def test_company_snapshot_does_not_mutate_on_get_and_defaults_revision_zero(ctx):
    ctx.settings.delete()
    before = AuditLog.objects.count()
    response = ctx.client.get("/payroll-settings/", **ctx.headers)
    assert response.status_code == 200 and response.data["data"]["id"] is None
    assert response.data["data"]["settings_revision"] == 0
    assert not PayrollSettings.objects.filter(company=ctx.company).exists() and AuditLog.objects.count() == before
    assert create(ctx).settings_snapshot["settings_revision"] == 0


def test_head_office_aggregate_reads_but_writes_are_blocked(ctx):
    office = OrganizationNode.objects.create(code="P2_HO", name="Head office", node_type="head_office")
    UserOrganizationAccess.objects.create(user=ctx.user, organization=office)
    run = create(ctx)
    headers = {"HTTP_X_ACTIVE_COMPANY_ID": str(office.pk)}
    assert (
        ctx.client.post(f"/payroll-runs/{run.pk}/finalize/", {"confirm": True}, format="json", **headers).status_code
        == 403
    )
    run.refresh_from_db()
    assert run.status == "DRAFT"


def test_cfo_role_isolated_allowlist_and_nonfinance_roles_denied(ctx):
    ctx.user.groups.clear()
    ctx.user.groups.add(Group.objects.get_or_create(name="CFO")[0])
    assert ctx.client.get("/payroll-settings/", **ctx.headers).status_code == 200
    for role in ("CEO", "Manager"):
        ctx.user.groups.clear()
        ctx.user.groups.add(Group.objects.get_or_create(name=role)[0])
        assert ctx.client.get("/payroll-settings/", **ctx.headers).status_code == 403


def test_publication_reconciliation_blocks_tampered_frozen_slip(ctx):
    run = create(ctx)
    finalize(ctx, run)
    slip = run.payslips.get()
    Payslip.objects.filter(pk=slip.pk).update(net_salary=Decimal("1"))
    response = ctx.client.post(f"/payroll-runs/{run.pk}/generate-payslips/", {}, format="json", **ctx.headers)
    assert response.status_code == 422 and response.data["errors"][0]["code"] == "payroll_reconciliation_failed"
    run.refresh_from_db()
    assert run.status == "COMPLETED"


def test_export_formula_prefixes_and_no_money_float(ctx):
    run = create(ctx)
    item = run.items.get()
    item.employee_name = '=HYPERLINK("https://example.invalid")'
    item.save(update_fields=["employee_name"])
    response = ctx.client.get(f"/payroll-runs/{run.pk}/export/", {"file_format": "xlsx"}, **ctx.headers)
    wb = openpyxl.load_workbook(BytesIO(response.content))
    assert wb["Payroll"].cell(2, 2).value.startswith("'=HYPERLINK")
    assert wb["Payroll"].cell(2, 2).data_type == "s"
    assert not any(cell.data_type == "f" for sheet in wb for row in sheet for cell in row)


def penalty(ctx, fine="0.00", wage="0.00", evidence=None):
    catalog = PenaltyCatalog.objects.create(
        code="P2",
        category="p2",
        recurrence_group="p2",
        title_en="Test penalty",
        title_ar="اختبار",
        count_period="monthly",
        source_page=1,
        source_row=1,
    )
    record = PenaltyRecord.objects.create(
        company=ctx.company,
        employee_profile=ctx.profile,
        catalog=catalog,
        occurred_on=date(2026, 9, 1),
        occurrence_number=1,
        action="deduction",
        amount=fine,
        extra_wage_amount=wage,
        total_deduction_amount=Decimal(fine) + Decimal(wage),
        status="issued",
        source="hr",
        source_kind="absence",
        evidence=evidence or {},
    )
    return PenaltyDeduction.objects.create(
        penalty=record,
        company=ctx.company,
        employee_profile=ctx.profile,
        intended_year=2026,
        intended_month=9,
        amount=record.total_deduction_amount,
        status="approved",
    )


def test_exact_article70_cap_then_rounded_audit_metadata(ctx):
    for field in (
        "accommodation_allowance",
        "transportation_allowance",
        "telephone_allowance",
        "petrol_allowance",
        "other_allowance",
    ):
        setattr(ctx.profile, field, Decimal(0))
    ctx.profile.basic_salary = ctx.profile.total_salary = Decimal("1000")
    ctx.profile.save()
    claim = penalty(ctx, fine="166.50")
    run = create(ctx)
    claim.refresh_from_db()
    assert (
        claim.status == "claimed"
        and claim.claimed_exact_amount == Decimal("166.50")
        and claim.claimed_amount == Decimal("167")
    )
    assert dto.item_dto(run.items.get())["penalty_deduction"] == "167.00"
    audit = AuditLog.objects.get(action="penalty_claimed_by_payroll", entity_id=str(claim.pk))
    assert audit.metadata["exact_amount"] == "166.50" and audit.metadata["rounded_amount"] == "167"
    finalize(ctx, run)
    claim.refresh_from_db()
    assert claim.status == "applied"


def test_penalty_wage_absence_overlap_and_cancel_releases_whole_claim(ctx):
    AttendanceRecord.objects.create(employee_profile=ctx.profile, date=date(2026, 9, 1), status="ABSENT")
    claim = penalty(ctx, wage="366.67", evidence={"wage_absence_dates": ["2026-09-01"]})
    run = create(ctx)
    value = dto.item_dto(run.items.get())
    assert value["absence_deduction"] == "0.00" and value["penalty_deduction"] == "367.00"
    response = ctx.client.post(
        f"/payroll-runs/{run.pk}/cancel/",
        {"expected_revision": run.revision, "reason": "Correction"},
        format="json",
        **ctx.headers,
    )
    assert response.status_code == 200, response.data
    claim.refresh_from_db()
    assert claim.status == "approved" and claim.payroll_run_id is None and claim.claimed_amount == 0


def test_unaffordable_penalty_stays_whole_pending_and_manual_reentry(ctx):
    claim = penalty(ctx, fine="500.00")
    run = create(ctx)
    value = edit(
        ctx, run, {"absence_days_override": "30.00", "zero_salary_reason": "full_absence", "other_deduction": "50.00"}
    )
    claim.refresh_from_db()
    assert claim.status == "approved" and claim.payroll_run_id is None and claim.claimed_amount == 0
    assert value["carry_forward"]["uncollected"] == "550.00"
    assert any(
        c["manual_reentry_required"] for c in value["carry_forward"]["by_category"] if c["type"] == "OTHER_DEDUCTION"
    )
    assert not PayrollCarryForward.objects.exists()


def test_archive_automatically_ends_assignment_with_actor_audit(ctx):
    obj = ctx.profile.payroll_accommodation_assignments.get()
    at = timezone.make_aware(datetime(2026, 9, 15, 12))
    retire_biotime_mapping_and_archive_profile(ctx.profile, {}, ctx.user, "Contract ended", at)
    obj.refresh_from_db()
    assert obj.end_date == date(2026, 9, 15)
    event = AuditLog.objects.get(action="payroll_accommodation_assignment_ended", entity_id=str(obj.pk))
    assert event.actor == ctx.user
    run = create(ctx)
    assert dto.item_dto(run.items.get())["salary_days"] == "15.00"


@pytest.mark.parametrize("mode,expected", [("FULL_MONTH", "350.00"), ("PRORATED_BY_DAYS", "175.00")])
def test_midmonth_provision_modes_and_full_absence_base(ctx, mode, expected):
    ctx.profile.hire_date = date(2026, 9, 16)
    ctx.profile.save(update_fields=["hire_date"])
    ctx.settings.provision_proration = mode
    ctx.settings.save()
    run = create(ctx)
    value = edit(ctx, run, {"absence_days_override": "2.00"})
    assert (
        value["gross_salary"] == "5500.00"
        and value["absence_deduction"] == "733.00"
        and value["net_salary"] == "4767.00"
    )
    assert value["leave_provision"] == expected and run.settings_snapshot["provision_proration"] == mode


def test_suspended_run_decision_and_future_hire_exclusion(ctx):
    ctx.profile.employment_status = "SUSPENDED"
    ctx.profile.save(update_fields=["employment_status"])
    run = create(ctx)
    response = finalize(ctx, run, 422)
    assert "suspended_pay_decision_required" in {e["code"] for e in response.data["errors"]}
    value = edit(ctx, run, {"suspended_pay_decision": "ZERO_PAY", "zero_salary_reason": "suspended"})
    assert value["gross_salary"] == "0.00" and value["net_salary"] == "0.00"
    finalize(ctx, run)


def test_identical_item_edit_has_no_effective_change_or_fabricated_audit(ctx):
    run = create(ctx)
    edit(ctx, run, {"bonus": "50.00"})
    run.refresh_from_db()
    revision, count = run.revision, AuditLog.objects.count()
    response = ctx.client.patch(
        f"/payroll-runs/{run.pk}/items/{run.items.get().pk}/",
        {"expected_revision": revision, "changes": {"bonus": "50.00"}, "reason": "Repeat correction"},
        format="json",
        **ctx.headers,
    )
    assert response.status_code == 422 and response.data["errors"][0]["code"] == "field_not_editable"
    run.refresh_from_db()
    assert run.revision == revision and AuditLog.objects.count() == count


@pytest.mark.parametrize("adjusted", [False, True])
def test_monthly_salary_inputs_midmonth_and_prefilled_patch_is_noop(ctx, adjusted):
    ctx.profile.hire_date = date(2026, 6, 16)
    ctx.profile.save(update_fields=["hire_date"])
    run = create(ctx, month=6)
    if adjusted:
        edit(ctx, run, {"salary_components": {"basic_salary": "6100.00"}})
    run.refresh_from_db()
    item = run.items.get()
    path = f"/payroll-runs/{run.pk}/items/{item.pk}/"
    value = ctx.client.get(path, **ctx.headers).data["data"]
    assert value["monthly_salary_components"] == {
        "basic_salary": "6100.00" if adjusted else "6000.00",
        "accommodation_allowance": "3000.00",
        "transportation_allowance": "1500.00",
        "telephone_allowance": "500.00",
        "petrol_allowance": "0.00",
        "other_allowance": "0.00",
    }
    assert value["salary_adjusted"] is adjusted
    assert value["salary_days"] == "15.00"
    assert value["basic_salary"] == ("3050.00" if adjusted else "3000.00")
    assert value["accommodation_allowance"] == "1500.00"
    assert value["transportation_allowance"] == "750.00"
    assert value["telephone_allowance"] == "250.00"
    assert value["full_monthly_gross"] == ("11100.00" if adjusted else "11000.00")
    assert value["gross_salary"] == ("5550.00" if adjusted else "5500.00")
    before = (run.revision, item.input_snapshot, AuditLog.objects.count(), PayrollAdjustment.objects.count())
    # A later profile edit must not change the run's editor defaults or create
    # an override when the client sends the snapshotted prefill unchanged.
    ctx.profile.basic_salary = Decimal("9999.00")
    ctx.profile.save(update_fields=["basic_salary"])
    assert ctx.client.get(path, **ctx.headers).data["data"] == value
    response = ctx.client.patch(
        path,
        {
            "expected_revision": run.revision,
            "changes": {"salary_components": value["monthly_salary_components"]},
            "reason": "Unchanged editor prefill",
        },
        format="json",
        **ctx.headers,
    )
    assert response.status_code == 422 and response.data["errors"][0]["code"] == "field_not_editable"
    run.refresh_from_db()
    item.refresh_from_db()
    assert (run.revision, item.input_snapshot, AuditLog.objects.count(), PayrollAdjustment.objects.count()) == before
    assert ctx.client.get(path, **ctx.headers).data["data"] == value


def test_monthly_salary_prefill_with_another_edit_does_not_activate_salary_adjustment(ctx):
    run = create(ctx, month=6)
    monthly = dto.item_dto(run.items.get())["monthly_salary_components"]
    value = edit(ctx, run, {"salary_components": monthly, "bonus": "50.00"})
    assert value["monthly_salary_components"] == monthly and value["salary_adjusted"] is False
    assert value["bonus"] == "50.00"
    assert not PayrollAdjustment.objects.filter(run=run, type="SALARY_ADJUSTMENT").exists()
    assert "salary_components" not in run.items.get().input_snapshot["overrides"]


def test_manual_inputs_preserve_entered_halalas_before_line_rounding(ctx):
    run = create(ctx)
    defaults = {
        "absence_days_override": None,
        "overtime_normal_hours": "0.00",
        "weekly_rest_work_days": "0.00",
        "official_holiday_work_days": "0.00",
        "bonus": "0.00",
        "incentive": "0.00",
        "other_addition": "0.00",
        "other_deduction": "0.00",
        "override_overtime_eligibility": False,
    }
    assert dto.item_dto(run.items.get())["manual_inputs"] == defaults
    entered = {
        **defaults,
        "absence_days_override": "2.00",
        "overtime_normal_hours": "1.25",
        "weekly_rest_work_days": "0.50",
        "official_holiday_work_days": "0.25",
        "bonus": "250.50",
        "incentive": "10.49",
        "other_addition": "20.50",
        "other_deduction": "1.50",
        "override_overtime_eligibility": True,
    }
    value = edit(ctx, run, entered)
    assert value["manual_inputs"] == entered
    assert value["bonus"] == "251.00"
    assert next(row["amount"] for row in value["additions"] if row["type"] == "BONUS") == "251.00"
    item = run.items.get()
    assert item.input_snapshot["overrides"]["bonus"] == "250.50"
    assert ctx.client.get(f"/payroll-runs/{run.pk}/items/{item.pk}/", **ctx.headers).data["data"] == value


def test_legacy_unknown_monthly_salary_is_not_invented_from_paid_or_profile_values(ctx):
    run = create(ctx)
    item = run.items.get()
    item.input_snapshot = {"schema_version": 0, "legacy_unknown": True}
    value = dto.item_dto(item)
    assert value["monthly_salary_components"] is None and value["salary_adjusted"] is False


def test_publish_notifies_once_after_commit_without_company_cost(ctx, django_capture_on_commit_callbacks):
    run = create(ctx)
    finalize(ctx, run)
    with patch("in_app_notifications.dispatcher.dispatch_notification_channels") as notify:
        with django_capture_on_commit_callbacks(execute=True):
            response = ctx.client.post(f"/payroll-runs/{run.pk}/generate-payslips/", {}, format="json", **ctx.headers)
        assert response.status_code == 200 and notify.call_count == 1
        values = notify.call_args.kwargs
        assert values["recipient"] == ctx.employee and values["company"] == ctx.company
        assert "total_company_cost" not in json.dumps(values.get("metadata", {}))
        with django_capture_on_commit_callbacks(execute=True):
            assert (
                ctx.client.post(
                    f"/payroll-runs/{run.pk}/generate-payslips/", {}, format="json", **ctx.headers
                ).status_code
                == 200
            )
        assert notify.call_count == 1


def test_carry_reservation_prevents_double_collection_and_cancel_releases(ctx):
    ctx.profile.deduct_gosi_from_employee = True
    ctx.profile.save(update_fields=["deduct_gosi_from_employee"])
    first = create(ctx)
    edit(ctx, first, {"absence_days_override": "30.00", "zero_salary_reason": "full_absence"})
    finalize(ctx, first)
    balance = PayrollCarryForward.objects.get(origin_run=first)
    october = create(ctx, 10)
    november = create(ctx, 11)
    assert dto.item_dto(november.items.get())["gosi_employee_deduction"] == "180.00"
    assert (
        ctx.client.post(
            f"/payroll-runs/{october.pk}/cancel/",
            {"expected_revision": october.revision, "reason": "Release reservation"},
            format="json",
            **ctx.headers,
        ).status_code
        == 200
    )
    balance.refresh_from_db()
    assert balance.remaining_amount == Decimal("180")
    engine.recalculate(november, request=ctx.request)
    assert dto.item_dto(november.items.get())["gosi_employee_deduction"] == "360.00"
    finalize(ctx, november)
    balance.refresh_from_db()
    assert balance.remaining_amount == 0


def test_assignment_service_overlap_history_and_soft_delete(ctx):
    obj = ctx.profile.payroll_accommodation_assignments.get()
    response = ctx.client.post(
        "/payroll-accommodation-assignments/",
        {
            "employee_profile_id": ctx.profile.pk,
            "status": "none",
            "rent_id": None,
            "start_date": "2026-09-01",
            "end_date": "2026-09-30",
            "reason": "Accommodation",
        },
        format="json",
        **ctx.headers,
    )
    assert response.status_code == 409, response.data
    response = ctx.client.patch(
        f"/payroll-accommodation-assignments/{obj.pk}/",
        {"end_date": "2026-09-15", "reason": "Move"},
        format="json",
        **ctx.headers,
    )
    assert response.status_code == 200, response.data
    history = ctx.client.get(f"/payroll-accommodation-assignments/{obj.pk}/history/", **ctx.headers).data["data"][
        "items"
    ][0]
    assert (
        history["type"] == "ASSIGNMENT_CHANGE" and history["assignment_id"] == obj.pk and "run_revision" not in history
    )
    assert (
        ctx.client.delete(
            f"/payroll-accommodation-assignments/{obj.pk}/",
            {"reason": "Withdraw assignment"},
            format="json",
            **ctx.headers,
        ).status_code
        == 204
    )
    obj.refresh_from_db()
    assert not obj.is_active


def test_unknown_cost_subtotals_and_medical_single_run_warning(ctx):
    ctx.profile.payroll_accommodation_assignments.all().delete()
    ctx.settings.no_medical_insurance = False
    ctx.settings.save()
    run = create(ctx)
    item = dto.item_dto(run.items.get())
    assert item["medical_provision"] is None and item["accommodation_cost"] is None and item["cost_incomplete"]
    summary = dto.summary(run)
    assert summary["provisions"]["medical"] == "0.00" and summary["unknown_provision_counts"]["MEDICAL"] == 1
    assert [w["code"] for w in dto.validation(run)["warnings"]].count("medical_invoice_missing") == 1


@pytest.mark.parametrize(
    "path",
    [
        "payroll-settings",
        "payroll-settings/salary-bases",
        "payroll-settings/enums",
        "payroll-settings/history",
        "payroll-settings/employees/1",
        "payroll-settings/employees/1/history",
        "payroll-runs/1/items/2",
        "payroll-runs/1/validation-report",
        "payroll-runs/1/recalculate",
        "payroll-runs/1/cancel",
        "payroll-runs/1/items/2/adjustments",
        "payroll-runs/1/history",
        "payroll-runs/1/accommodation-report",
        "payroll-accommodation-assignments/1/history",
    ],
)
def test_approved_routes_accept_both_slash_forms(path):
    assert resolve("/" + path).func == resolve("/" + path + "/").func


def test_arabic_enums_and_private_headers(ctx):
    response = ctx.client.get("/payroll-settings/enums/", HTTP_ACCEPT_LANGUAGE="ar", **ctx.headers)
    assert response.status_code == 200
    codes = response.data["data"]["validation_codes"]
    assert len(codes) == 44 and any("أ" in r["label"] or "ا" in r["label"] for r in codes)
    assert response["Cache-Control"] == "private, no-store"
