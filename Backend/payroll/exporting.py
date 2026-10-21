"""Read-only HR exports. Exact money text and escaped cells preserve snapshots."""

import json
import re
from datetime import date
from decimal import Decimal
from io import BytesIO

import openpyxl

from . import dto
from .models import PayrollAdjustment

LEGACY_HEADERS = [
    "employee_id",
    "employee_name",
    "department",
    "position",
    "basic_salary",
    "allowances",
    "deductions",
    "net_salary",
]
EXTRA_FIELDS = [
    "employee_profile_id",
    "salary_days",
    "full_monthly_gross",
    "gross_salary",
    *dto.SCALARS,
    "penalty_deduction",
    "total_additions",
    "total_provisions",
    "total_company_cost",
    "cost_incomplete",
    "accommodation_status",
    "zero_salary_reason",
    "zero_salary_reason_detail",
    "suspended_pay_decision",
]


def safe_cell(value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True)
    if isinstance(value, Decimal):
        value = dto.amount(value)
    if (
        isinstance(value, str)
        and not re.fullmatch(r"-?\d+(?:\.\d+)?", value)
        and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r"))
    ):
        return "'" + value
    return value


def payroll_rows(run, items):
    headers = (
        LEGACY_HEADERS
        + EXTRA_FIELDS
        + [
            "period_start",
            "period_end",
            "calendar_days",
            "carry_forward",
            "settings_revision",
            "validation_errors",
            "validation_warnings",
        ]
    )
    counts = dto.validation(run)["counts"]
    run_value = dto.run_dto(run)
    effective_calendar_days = run_value["calendar_days"] or (
        date.fromisoformat(run_value["period_end"]) - date.fromisoformat(run_value["period_start"])
    ).days + 1
    rows = [headers]
    for item in items:
        value = dto.item_dto(item)
        legacy = [
            item.employee_id,
            item.employee_name,
            item.department,
            item.position,
            dto.amount(item.basic_salary),
            dto.amount(item.total_allowances),
            dto.amount(item.total_deductions),
            dto.amount(item.net_salary),
        ]
        rows.append(
            [
                safe_cell(v)
                for v in legacy
                + [value[k] for k in EXTRA_FIELDS]
                + [
                    run_value["period_start"],
                    run_value["period_end"],
                    effective_calendar_days,
                    value["carry_forward"],
                    run.settings_snapshot.get("settings_revision", 0),
                    counts["errors"],
                    counts["warnings"],
                ]
            ]
        )
    return rows


def workbook(run, items):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Payroll"
    rows = payroll_rows(run, items)
    rows[0][:8] = [
        "Employee ID",
        "Employee Name",
        "Department",
        "Position",
        "Basic Salary",
        "Allowances",
        "Deductions",
        "Net Salary",
    ]
    for row in rows:
        ws.append(row)
    for name, category in (("Lines", None), ("Provisions", "PROVISION")):
        sheet = wb.create_sheet(name)
        sheet.append(
            [
                "item_id",
                "employee_id",
                "category",
                "type",
                "label",
                "amount",
                "quantity",
                "requested_amount",
                "collected_amount",
                "uncollected_amount",
                "carried_in",
                "source_period",
                "source_kind",
                "source_id",
            ]
        )
        for item in items:
            for line in item.lines.order_by("id"):
                if category and line.category != category:
                    continue
                value = dto.line_dto(line)
                sheet.append(
                    [
                        safe_cell(v)
                        for v in [item.pk, item.employee_id, line.category]
                        + [
                            value[k]
                            for k in (
                                "type",
                                "label",
                                "amount",
                                "quantity",
                                "requested_amount",
                                "collected_amount",
                                "uncollected_amount",
                                "carried_in",
                                "source_period",
                                "source_kind",
                                "source_id",
                            )
                        ]
                    ]
                )
    report = dto.validation(run)
    sheet = wb.create_sheet("Validation")
    sheet.append(["code", "severity", "message", "employee", "params"])
    for issue in report["errors"] + report["warnings"]:
        sheet.append([safe_cell(issue.get(k)) for k in ("code", "severity", "message", "employee", "params")])
    sheet = wb.create_sheet("Adjustments")
    columns = [
        "id",
        "audit_log_id",
        "type",
        "changed_fields",
        "original_value",
        "new_value",
        "reason",
        "actor",
        "created_at",
        "run_revision",
    ]
    sheet.append(columns)
    for adjustment in PayrollAdjustment.objects.filter(run=run).select_related("actor").order_by("id"):
        value = dto.adjustment_dto(adjustment)
        sheet.append([safe_cell(value.get(k)) for k in columns])
    sheet = wb.create_sheet("Accommodation")
    sheet.append(
        [
            "rent_id",
            "rent_name",
            "kind",
            "employee",
            "in_payroll",
            "rounded_share",
            "raw_decimal_approximation",
            "raw_numerator",
            "raw_denominator",
        ]
    )
    for allocation in run.rent_allocations.order_by("rent_id"):
        value = dto.rent_dto(allocation)
        for resident in value["residents"]:
            raw = resident["raw_share"]
            sheet.append(
                [
                    safe_cell(v)
                    for v in [
                        value["rent_id"],
                        value["rent_name"],
                        "resident",
                        resident["employee"],
                        resident["in_payroll"],
                        resident["rounded_share"],
                        raw["decimal"],
                        raw["numerator"],
                        raw["denominator"],
                    ]
                ]
            )
        for kind in (
            "monthly_rent",
            "period_rent",
            "vacant_unallocated_raw",
            "not_in_payroll_raw",
            "rounding_difference",
        ):
            raw = value[kind]
            sheet.append(
                [
                    safe_cell(v)
                    for v in [
                        value["rent_id"],
                        value["rent_name"],
                        kind,
                        None,
                        None,
                        None,
                        raw["decimal"],
                        raw["numerator"],
                        raw["denominator"],
                    ]
                ]
            )
    for sheet in wb:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
